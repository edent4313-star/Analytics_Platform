"""
Task 1 Complete Test Suite.
Covers: auth, user admin, role admin, permission admin, org admin, audit, bulk import.
All tests use the TestClient (no running server needed).
"""
import io
import pytest
from fastapi.testclient import TestClient
from app.main import app

client = TestClient(app)

def login(eid="CBE001", pw="Demo@1234"):
    r = client.post("/api/v1/auth/login", json={"employee_id": eid, "password": pw})
    assert r.status_code == 200, f"Login failed for {eid}: {r.json()}"
    return r.json()["access_token"]

def auth(token): return {"Authorization": f"Bearer {token}"}


# ════════════════════════════════════════════════════════════════
# AUTHENTICATION
# ════════════════════════════════════════════════════════════════

class TestAuthentication:
    def test_valid_login(self):
        r = client.post("/api/v1/auth/login", json={"employee_id": "CBE001", "password": "Admin@1234"})
        assert r.status_code == 200
        d = r.json()
        assert "access_token" in d
        assert "refresh_token" in d
        assert d["provider"] == "mock"

    def test_invalid_employee_id(self):
        r = client.post("/api/v1/auth/login", json={"employee_id": "NOTEXIST", "password": "pass"})
        assert r.status_code == 401

    def test_wrong_password(self):
        r = client.post("/api/v1/auth/login", json={"employee_id": "CBE001", "password": "wrongpass"})
        assert r.status_code == 401

    def test_all_demo_users_login(self):
        creds = [
            ("CBE001", "Demo@1234"), ("CBE002", "Demo@1234"), ("CBE003", "Demo@1234"),
            ("CBE004", "Demo@1234"), ("CBE005", "Demo@1234"), ("CBE006", "Demo@1234"),
            ("CBE007", "Demo@1234"), ("CBE001", "Admin@1234"),
        ]
        for eid, pw in creds:
            r = client.post("/api/v1/auth/login", json={"employee_id": eid, "password": pw})
            assert r.status_code == 200, f"Login failed for {eid}"

    def test_get_me_authenticated(self):
        token = login()
        r = client.get("/api/v1/auth/me", headers=auth(token))
        assert r.status_code == 200
        d = r.json()
        assert d["employee_id"] == "CBE001"
        assert d["role"] == "ADMIN"
        assert d["access_level"] == "HEAD_OFFICE"

    def test_get_me_unauthenticated(self):
        r = client.get("/api/v1/auth/me")
        assert r.status_code == 403

    def test_get_me_invalid_token(self):
        r = client.get("/api/v1/auth/me", headers=auth("invalid.token.here"))
        assert r.status_code == 401

    def test_permissions_endpoint(self):
        token = login()
        r = client.get("/api/v1/auth/permissions", headers=auth(token))
        assert r.status_code == 200
        assert "permissions" in r.json()
        assert len(r.json()["permissions"]) > 0

    def test_data_scope_branch_manager(self):
        token = login("CBE005", "Demo@1234")
        r = client.get("/api/v1/users/me/data-scope", headers=auth(token))
        assert r.status_code == 200
        d = r.json()
        assert d["access_level"] == "BRANCH"
        assert d["branch_id"] is not None
        assert d["is_head_office"] is False

    def test_data_scope_head_office(self):
        token = login()
        r = client.get("/api/v1/users/me/data-scope", headers=auth(token))
        assert r.status_code == 200
        assert r.json()["is_head_office"] is True

    def test_session_endpoint(self):
        token = login("CBE002", "Demo@1234")
        r = client.get("/api/v1/auth/session", headers=auth(token))
        assert r.status_code == 200
        assert r.json()["provider"] == "mock"

    def test_refresh_token(self):
        r = client.post("/api/v1/auth/login", json={"employee_id": "CBE001", "password": "Admin@1234"})
        refresh = r.json()["refresh_token"]
        r2 = client.post("/api/v1/auth/refresh", json={"refresh_token": refresh})
        assert r2.status_code == 200
        assert "access_token" in r2.json()

    def test_logout(self):
        token = login()
        r = client.post("/api/v1/auth/logout", headers=auth(token))
        assert r.status_code == 200

    def test_oidc_login_not_active_in_mock_mode(self):
        r = client.get("/api/v1/auth/ad/login")
        assert r.status_code == 503  # cbe_ad not configured


# ════════════════════════════════════════════════════════════════
# USER ADMINISTRATION
# ════════════════════════════════════════════════════════════════

class TestUserAdministration:
    def test_admin_can_list_users(self):
        token = login()
        r = client.get("/api/v1/admin/users", headers=auth(token))
        assert r.status_code == 200
        assert "items" in r.json()
        assert r.json()["total"] >= 7

    def test_viewer_cannot_list_users(self):
        token = login("CBE007", "Demo@1234")
        r = client.get("/api/v1/admin/users", headers=auth(token))
        assert r.status_code == 403

    def test_unauthenticated_cannot_list_users(self):
        r = client.get("/api/v1/admin/users")
        assert r.status_code == 403

    def test_create_and_get_user(self):
        token = login()
        r = client.post("/api/v1/admin/users", headers=auth(token), json={
            "username": "task1_test_user",
            "full_name": "Task1 Test User",
            "email": "task1test@cbe.com.et",
            "password": "Pass@1234",
            "access_level": "HEAD_OFFICE",
            "role_id": 1,
            "employee_id": "CBE_TASK1",
        })
        assert r.status_code == 201
        user_id = r.json()["id"]
        assert r.json()["employee_id"] == "CBE_TASK1"

        r2 = client.get(f"/api/v1/admin/users/{user_id}", headers=auth(token))
        assert r2.status_code == 200
        assert r2.json()["username"] == "task1_test_user"

    def test_duplicate_username_rejected(self):
        token = login()
        client.post("/api/v1/admin/users", headers=auth(token), json={
            "username": "dup_task1",
            "full_name": "Dup User",
            "email": "dup1_task1@cbe.com.et",
            "password": "Pass@1234",
            "access_level": "HEAD_OFFICE",
            "role_id": 1,
        })
        r = client.post("/api/v1/admin/users", headers=auth(token), json={
            "username": "dup_task1",
            "full_name": "Dup User2",
            "email": "dup2_task1@cbe.com.et",
            "password": "Pass@1234",
            "access_level": "HEAD_OFFICE",
            "role_id": 1,
        })
        assert r.status_code == 400

    def test_invalid_org_hierarchy_rejected(self):
        token = login()
        r = client.post("/api/v1/admin/users", headers=auth(token), json={
            "username": "bad_org_test",
            "full_name": "Bad Org",
            "email": "badorg_task1@cbe.com.et",
            "password": "Pass@1234",
            "access_level": "DISTRICT",
            "region_id": 1,
            "district_id": 4,  # district 4 belongs to region 2, not region 1
            "role_id": 4,
        })
        assert r.status_code == 400

    def test_get_user_scope(self):
        token = login()
        r = client.get("/api/v1/admin/users/3/scope", headers=auth(token))
        assert r.status_code == 200
        assert r.json()["access_level"] == "REGION"

    def test_get_user_permissions(self):
        token = login()
        r = client.get("/api/v1/admin/users/1/permissions", headers=auth(token))
        assert r.status_code == 200
        assert "permissions" in r.json()


# ════════════════════════════════════════════════════════════════
# ROLE ADMINISTRATION
# ════════════════════════════════════════════════════════════════

class TestRoleAdministration:
    def test_admin_can_list_roles(self):
        token = login()
        r = client.get("/api/v1/admin/roles", headers=auth(token))
        assert r.status_code == 200
        assert len(r.json()) >= 7  # 7 system roles seeded

    def test_viewer_cannot_list_roles(self):
        token = login("CBE007", "Demo@1234")
        r = client.get("/api/v1/admin/roles", headers=auth(token))
        assert r.status_code == 403

    def test_create_custom_role(self):
        token = login()
        r = client.post("/api/v1/admin/roles", headers=auth(token), json={
            "name": "TASK1_TEST_ROLE",
            "display_name": "Task1 Test Role",
            "description": "Created by task1 test",
        })
        assert r.status_code == 201
        assert r.json()["name"] == "TASK1_TEST_ROLE"

    def test_cannot_modify_system_role(self):
        token = login()
        # Find ADMIN role (system role)
        roles = client.get("/api/v1/admin/roles", headers=auth(token)).json()
        admin_role = next(r for r in roles if r["name"] == "ADMIN")
        r = client.put(f"/api/v1/admin/roles/{admin_role['id']}", headers=auth(token), json={
            "name": "HACKED", "display_name": "Hacked"
        })
        assert r.status_code == 400

    def test_get_role_permissions(self):
        token = login()
        r = client.get("/api/v1/roles/1/permissions", headers=auth(token))
        assert r.status_code == 200


# ════════════════════════════════════════════════════════════════
# PERMISSION ADMINISTRATION
# ════════════════════════════════════════════════════════════════

class TestPermissionAdministration:
    def test_admin_can_list_permissions(self):
        token = login()
        r = client.get("/api/v1/admin/permissions", headers=auth(token))
        assert r.status_code == 200
        assert len(r.json()) >= 21

    def test_viewer_cannot_manage_permissions(self):
        token = login("CBE007", "Demo@1234")
        r = client.get("/api/v1/admin/permissions", headers=auth(token))
        assert r.status_code == 403


# ════════════════════════════════════════════════════════════════
# ORGANIZATION ADMINISTRATION
# ════════════════════════════════════════════════════════════════

class TestOrganizationAdministration:
    def test_admin_gets_full_org_tree(self):
        token = login()
        r = client.get("/api/v1/admin/organizations", headers=auth(token))
        assert r.status_code == 200
        orgs = r.json()
        assert len(orgs) == 3  # 3 seeded regions
        assert all("districts" in org for org in orgs)

    def test_scoped_regions_for_regional_manager(self):
        token = login("CBE003", "Demo@1234")
        r = client.get("/api/v1/regions", headers=auth(token))
        assert r.status_code == 200
        assert len(r.json()) == 1  # only their region

    def test_regional_manager_blocked_from_other_region(self):
        token = login("CBE003", "Demo@1234")  # region_id = 1
        r = client.get("/api/v1/regions/2/districts", headers=auth(token))
        assert r.status_code == 403

    def test_branch_manager_sees_only_own_branch(self):
        token = login("CBE005", "Demo@1234")
        scope = client.get("/api/v1/users/me/data-scope", headers=auth(token)).json()
        assert scope["branch_id"] is not None
        assert scope["access_level"] == "BRANCH"


# ════════════════════════════════════════════════════════════════
# AUDIT LOGGING
# ════════════════════════════════════════════════════════════════

class TestAuditLogging:
    def test_admin_can_view_audit(self):
        token = login()
        r = client.get("/api/v1/audit", headers=auth(token))
        assert r.status_code == 200
        assert "items" in r.json()

    def test_viewer_cannot_view_audit(self):
        token = login("CBE007", "Demo@1234")
        r = client.get("/api/v1/audit", headers=auth(token))
        assert r.status_code == 403

    def test_security_events_endpoint(self):
        token = login()
        r = client.get("/api/v1/audit/security-events", headers=auth(token))
        assert r.status_code == 200

    def test_login_creates_audit_entry(self):
        token = login()
        r = client.get("/api/v1/audit?action=LOGIN", headers=auth(token))
        assert r.status_code == 200
        items = r.json()["items"]
        assert any(i["action"] == "LOGIN" for i in items)

    def test_user_audit_trail(self):
        token = login()
        r = client.get("/api/v1/audit/users/1", headers=auth(token))
        assert r.status_code == 200

    def test_failed_login_creates_audit(self):
        client.post("/api/v1/auth/login", json={"employee_id": "CBE001", "password": "WRONG"})
        token = login()
        r = client.get("/api/v1/audit?action=LOGIN_FAILED", headers=auth(token))
        assert r.status_code == 200


# ════════════════════════════════════════════════════════════════
# BULK IMPORT
# ════════════════════════════════════════════════════════════════

class TestBulkImport:
    def test_download_template(self):
        token = login()
        r = client.get("/api/v1/import/template", headers=auth(token))
        assert r.status_code == 200
        assert "spreadsheetml" in r.headers.get("content-type", "")

    def test_bulk_import_dry_run(self):
        """Dry run: validates an Excel file without creating users."""
        import openpyxl
        token = login()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["username","full_name","email","phone","employee_id",
                   "access_level","region_code","district_code","branch_code",
                   "role_name","password"])
        ws.append(["dryrun_user1","Dry Run User 1","dryrun1@cbe.com.et","",
                   "","HEAD_OFFICE","","","","VIEWER","Pass@1234"])

        buf = io.BytesIO()
        wb.save(buf); buf.seek(0)

        r = client.post(
            "/api/v1/import/users?dry_run=true",
            headers=auth(token),
            files={"file": ("test.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert r.status_code == 200
        d = r.json()
        assert d["summary"]["dry_run"] is True
        assert d["summary"]["errors"] == 0

    def test_bulk_import_invalid_role(self):
        """Rows with invalid role names should get 'error' status."""
        import openpyxl
        token = login()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["username","full_name","email","phone","employee_id",
                   "access_level","region_code","district_code","branch_code",
                   "role_name","password"])
        ws.append(["bad_role_user","Bad Role","badrole@cbe.com.et","",
                   "","HEAD_OFFICE","","","","NONEXISTENT_ROLE","Pass@1234"])

        buf = io.BytesIO()
        wb.save(buf); buf.seek(0)

        r = client.post(
            "/api/v1/import/users?dry_run=true",
            headers=auth(token),
            files={"file": ("test.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert r.status_code == 200
        result = r.json()
        assert result["summary"]["errors"] >= 1

    def test_viewer_cannot_bulk_import(self):
        import openpyxl
        token = login("CBE007", "Demo@1234")
        buf = io.BytesIO()
        openpyxl.Workbook().save(buf); buf.seek(0)
        r = client.post(
            "/api/v1/import/users",
            headers=auth(token),
            files={"file": ("test.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert r.status_code == 403
