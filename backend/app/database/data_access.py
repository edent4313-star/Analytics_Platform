"""
Data Access Layer abstraction.
Provides a unified interface over PostgreSQL, Oracle, and Internal REST API
data sources. All analytical query execution goes through here.

This keeps the connector logic isolated — adding Oracle or a new API source
requires only a new connector class, not changes to the query engine.
"""
from __future__ import annotations

import logging
import re
from abc import ABC, abstractmethod
from typing import Any, Optional
from dataclasses import dataclass, field

logger = logging.getLogger(__name__)

_IDENT_RE = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")
_ALLOWED_AGGS = frozenset({"SUM", "COUNT", "AVG", "MIN", "MAX", "COUNT_DISTINCT"})


def quote_ident(name: str) -> str:
    if not name or not _IDENT_RE.match(name):
        raise ValueError("Invalid SQL identifier")
    return f'"{name}"'


def quote_table(schema_name: Optional[str], object_name: str) -> str:
    obj = quote_ident(object_name)
    if schema_name:
        return f"{quote_ident(schema_name)}.{obj}"
    return obj


def safe_agg(agg: Optional[str]) -> str:
    value = (agg or "SUM").upper()
    if value not in _ALLOWED_AGGS:
        raise ValueError("Invalid aggregation")
    return value


# ── Query abstraction ─────────────────────────────────────────────────────────

@dataclass
class FieldRef:
    """Reference to a field within a dataset."""
    field_name: str
    display_name: str = ""
    aggregation: Optional[str] = None  # SUM, COUNT, AVG, MIN, MAX, COUNT_DISTINCT


@dataclass
class DataScope:
    """
    Represents the organizational data scope for a user.
    Built by security/data_scope.py from the authenticated user's profile.
    NEVER constructed from user-supplied request parameters.

    Empty region/district/branch lists are NOT unrestricted.
    unrestricted_org_access must be set by an explicit authorization policy.
    """
    access_level: str  # HEAD_OFFICE | REGION | DISTRICT | BRANCH
    region_ids: Optional[list[int]] = None
    district_ids: Optional[list[int]] = None
    branch_ids: Optional[list[int]] = None
    unrestricted_org_access: bool = False
    deny_all: bool = False

    @property
    def is_unrestricted(self) -> bool:
        return bool(self.unrestricted_org_access) and not self.deny_all


def empty_result(page: int, page_size: int) -> "QueryResult":
    return QueryResult(columns=[], rows=[], total_count=0, page=page, page_size=page_size)


def apply_org_scope_clauses(
    qd: "QueryDefinition",
    scope: DataScope,
    params: dict,
    prefix: str = "",
) -> list[str]:
    """
    Return parameterized WHERE fragments for org scope.
    Fail closed: if the user is restricted and no matching scope column exists,
    return a clause that matches no rows.
    """
    if scope.deny_all:
        return ["1 = 0"]
    if scope.is_unrestricted:
        return []

    clauses: list[str] = []
    if scope.branch_ids and qd.branch_column:
        col = quote_ident(qd.branch_column)
        keys = [f"{prefix}branch_id_{i}" for i, _ in enumerate(scope.branch_ids)]
        clauses.append(f'{col} IN ({", ".join(":" + k for k in keys)})')
        for k, bid in zip(keys, scope.branch_ids):
            params[k] = bid
        return clauses
    if scope.district_ids and qd.district_column:
        col = quote_ident(qd.district_column)
        keys = [f"{prefix}district_id_{i}" for i, _ in enumerate(scope.district_ids)]
        clauses.append(f'{col} IN ({", ".join(":" + k for k in keys)})')
        for k, did in zip(keys, scope.district_ids):
            params[k] = did
        return clauses
    if scope.region_ids and qd.region_column:
        col = quote_ident(qd.region_column)
        keys = [f"{prefix}region_id_{i}" for i, _ in enumerate(scope.region_ids)]
        clauses.append(f'{col} IN ({", ".join(":" + k for k in keys)})')
        for k, rid in zip(keys, scope.region_ids):
            params[k] = rid
        return clauses

    # Restricted user but dataset has no applicable scope column → no rows
    return ["1 = 0"]


@dataclass
class QueryDefinition:
    """
    Defines what data a widget needs. Built by the query engine from
    the widget's config_json. All field references are validated against
    dataset_fields before execution.
    """
    dataset_id: int
    source_type: str                  # POSTGRESQL | ORACLE | INTERNAL_API
    connection_config: dict           # decrypted at query-engine level
    schema_name: Optional[str]
    object_name: str                  # table or view name
    dimensions: list[FieldRef] = field(default_factory=list)
    metrics: list[FieldRef] = field(default_factory=list)
    filters: list[dict] = field(default_factory=list)
    sort_by: Optional[str] = None
    sort_order: str = "DESC"
    limit: Optional[int] = None
    offset: int = 0
    # Org-scope column names in the target table
    region_column: Optional[str] = None
    district_column: Optional[str] = None
    branch_column: Optional[str] = None


@dataclass
class QueryResult:
    columns: list[str]
    rows: list[dict]
    total_count: int
    page: int = 1
    page_size: int = 25


# ── Connector interface ───────────────────────────────────────────────────────

class BaseConnector(ABC):
    """Abstract base for all data source connectors."""

    @abstractmethod
    def test_connection(self) -> bool:
        """Verify connectivity to the data source."""

    @abstractmethod
    def get_schema_objects(self, schema: str) -> list[str]:
        """Return list of tables/views available in the schema."""

    @abstractmethod
    def get_field_metadata(self, schema: str, object_name: str) -> list[dict]:
        """Return column metadata for a table/view."""

    @abstractmethod
    def execute_query(
        self,
        query_def: QueryDefinition,
        scope: DataScope,
        page: int = 1,
        page_size: int = 25,
    ) -> QueryResult:
        """Execute a query, applying organizational scope filters."""


# ── PostgreSQL connector ──────────────────────────────────────────────────────

class PostgreSQLConnector(BaseConnector):
    """
    Connects to an external analytical PostgreSQL database/schema.
    (Not the application DB — that uses SQLAlchemy ORM directly.)
    """

    def __init__(self, host: str, port: int, database: str,
                 username: str, password: str, schema: str = "public"):
        self.host = host
        self.port = port
        self.database = database
        self.username = username
        self.password = password
        self.schema = schema
        self._engine = None

    def _get_engine(self):
        if self._engine is None:
            from sqlalchemy import create_engine
            url = (
                f"postgresql://{self.username}:{self.password}"
                f"@{self.host}:{self.port}/{self.database}"
            )
            self._engine = create_engine(url, pool_pre_ping=True, pool_size=2)
        return self._engine

    def test_connection(self) -> bool:
        try:
            from sqlalchemy import text
            with self._get_engine().connect() as conn:
                conn.execute(text("SELECT 1"))
            return True
        except Exception as e:
            logger.error(f"PostgreSQL connection test failed: {e}")
            return False

    def get_schema_objects(self, schema: str) -> list[str]:
        from sqlalchemy import text, inspect
        try:
            insp = inspect(self._get_engine())
            return insp.get_table_names(schema=schema) + insp.get_view_names(schema=schema)
        except Exception as e:
            logger.error(f"Failed to get schema objects: {e}")
            return []

    def get_field_metadata(self, schema: str, object_name: str) -> list[dict]:
        from sqlalchemy import text, inspect
        try:
            insp = inspect(self._get_engine())
            columns = insp.get_columns(object_name, schema=schema)
            return [
                {
                    "field_name": col["name"],
                    "data_type": str(col["type"]),
                    "nullable": col.get("nullable", True),
                }
                for col in columns
            ]
        except Exception as e:
            logger.error(f"Failed to get field metadata: {e}")
            return []

    def execute_query(
        self,
        query_def: QueryDefinition,
        scope: DataScope,
        page: int = 1,
        page_size: int = 25,
    ) -> QueryResult:
        if scope.deny_all:
            return empty_result(page, page_size)
        from sqlalchemy import text
        sql, params = self._build_sql(query_def, scope, page, page_size)
        count_sql, count_params = self._build_count_sql(query_def, scope)
        try:
            with self._get_engine().connect() as conn:
                total = conn.execute(text(count_sql), count_params).scalar() or 0
                result = conn.execute(text(sql), params)
                rows = [dict(row._mapping) for row in result]
            return QueryResult(
                columns=list(rows[0].keys()) if rows else [],
                rows=rows,
                total_count=total,
                page=page,
                page_size=page_size,
            )
        except Exception as e:
            logger.error(f"Query execution failed: {e}")
            raise

    def _build_sql(
        self,
        qd: QueryDefinition,
        scope: DataScope,
        page: int,
        page_size: int,
    ) -> tuple[str, dict]:
        """
        Build a parameterized SELECT statement from QueryDefinition.
        Scope filters are injected as WHERE conditions — never trusted from
        user-supplied parameters.
        """
        table = quote_table(qd.schema_name, qd.object_name)
        select_parts = []
        params: dict = {}

        # Dimensions
        for dim in qd.dimensions:
            select_parts.append(quote_ident(dim.field_name))

        # Metrics with aggregation
        for i, metric in enumerate(qd.metrics):
            agg = safe_agg(metric.aggregation)
            field = quote_ident(metric.field_name)
            alias = quote_ident(metric.field_name)
            if agg == "COUNT_DISTINCT":
                select_parts.append(f"COUNT(DISTINCT {field}) AS {alias}")
            else:
                select_parts.append(f"{agg}({field}) AS {alias}")

        if not select_parts:
            select_parts = ["*"]

        select_clause = ", ".join(select_parts)
        where_clauses = apply_org_scope_clauses(qd, scope, params)

        # User-applied filters (already validated against dataset fields)
        for j, f in enumerate(qd.filters):
            col = f.get("field")
            op = f.get("operator", "eq")
            val = f.get("value")
            if col and val is not None:
                quoted = quote_ident(col)
                key = f"filter_{j}"
                if op == "eq":
                    where_clauses.append(f"{quoted} = :{key}")
                elif op == "gte":
                    where_clauses.append(f"{quoted} >= :{key}")
                elif op == "lte":
                    where_clauses.append(f"{quoted} <= :{key}")
                elif op == "like":
                    where_clauses.append(f"{quoted} ILIKE :{key}")
                    val = f"%{val}%"
                elif op == "in":
                    keys = [f"{key}_{k}" for k in range(len(val))]
                    where_clauses.append(f'{quoted} IN ({", ".join(":" + k for k in keys)})')
                    for k, v in zip(keys, val):
                        params[k] = v
                    continue
                params[key] = val

        where_sql = f"WHERE {' AND '.join(where_clauses)}" if where_clauses else ""

        group_by = ""
        if qd.dimensions and qd.metrics:
            dims = ", ".join(quote_ident(d.field_name) for d in qd.dimensions)
            group_by = f"GROUP BY {dims}"

        order_by = ""
        if qd.sort_by:
            direction = "DESC" if qd.sort_order.upper() == "DESC" else "ASC"
            order_by = f"ORDER BY {quote_ident(qd.sort_by)} {direction}"

        offset = (page - 1) * page_size
        sql = (
            f"SELECT {select_clause} FROM {table} "
            f"{where_sql} {group_by} {order_by} "
            f"LIMIT :limit OFFSET :offset"
        )
        params["limit"] = page_size
        params["offset"] = offset
        return sql.strip(), params

    def _build_count_sql(self, qd: QueryDefinition, scope: DataScope) -> tuple[str, dict]:
        table = quote_table(qd.schema_name, qd.object_name)
        # Build a simplified count query reusing the same where logic
        # For now wrap in subquery approach
        inner_sql, params = self._build_sql(qd, scope, page=1, page_size=1_000_000)
        # Remove LIMIT/OFFSET for counting
        base = inner_sql.rsplit("LIMIT", 1)[0].strip()
        count_sql = f"SELECT COUNT(*) FROM ({base}) AS _count_subquery"
        # Remove limit/offset from params
        params.pop("limit", None)
        params.pop("offset", None)
        return count_sql, params


# ── Oracle connector ──────────────────────────────────────────────────────────

class OracleConnector(BaseConnector):
    """
    Connects to an Oracle analytical database using python-oracledb thin mode.
    Thin mode requires no Oracle Instant Client installation.
    oracledb v26.0.0 is already installed in this environment.
    """

    def __init__(self, host: str, port: int, service_name: str,
                 username: str, password: str):
        self.host = host
        self.port = port
        self.service_name = service_name
        self.username = username
        self.password = password

    def _get_connection(self):
        import oracledb
        return oracledb.connect(
            user=self.username,
            password=self.password,
            dsn=f"{self.host}:{self.port}/{self.service_name}",
        )

    def test_connection(self) -> bool:
        try:
            conn = self._get_connection()
            conn.close()
            return True
        except Exception as e:
            logger.error(f"Oracle connection test failed: {e}")
            return False

    def get_schema_objects(self, schema: str) -> list[str]:
        try:
            with self._get_connection() as conn:
                cursor = conn.cursor()
                cursor.execute(
                    "SELECT object_name FROM all_objects "
                    "WHERE owner = :schema AND object_type IN ('TABLE','VIEW')",
                    schema=schema.upper(),
                )
                return [row[0] for row in cursor.fetchall()]
        except Exception as e:
            logger.error(f"Oracle get_schema_objects failed: {e}")
            return []

    def get_field_metadata(self, schema: str, object_name: str) -> list[dict]:
        try:
            with self._get_connection() as conn:
                cursor = conn.cursor()
                cursor.execute(
                    "SELECT column_name, data_type, nullable "
                    "FROM all_tab_columns "
                    "WHERE owner = :schema AND table_name = :obj "
                    "ORDER BY column_id",
                    schema=schema.upper(),
                    obj=object_name.upper(),
                )
                return [
                    {"field_name": r[0], "data_type": r[1], "nullable": r[2] == "Y"}
                    for r in cursor.fetchall()
                ]
        except Exception as e:
            logger.error(f"Oracle get_field_metadata failed: {e}")
            return []

    def execute_query(
        self,
        query_def: QueryDefinition,
        scope: DataScope,
        page: int = 1,
        page_size: int = 25,
    ) -> QueryResult:
        """
        Oracle uses bind variables with :name syntax — same as SQLAlchemy text().
        Pagination uses OFFSET/FETCH (Oracle 12c+).
        """
        if scope.deny_all:
            return empty_result(page, page_size)
        sql, params = self._build_oracle_sql(query_def, scope, page, page_size)
        try:
            with self._get_connection() as conn:
                cursor = conn.cursor()
                # Count query
                base_sql, base_params = self._build_oracle_sql(
                    query_def, scope, page=1, page_size=1_000_000
                )
                count_sql = f"SELECT COUNT(*) FROM ({base_sql.rsplit('OFFSET', 1)[0]}) t"
                cursor.execute(count_sql, base_params)
                total = cursor.fetchone()[0]

                cursor.execute(sql, params)
                columns = [col[0].lower() for col in cursor.description]
                rows = [dict(zip(columns, row)) for row in cursor.fetchall()]

            return QueryResult(
                columns=columns,
                rows=rows,
                total_count=total,
                page=page,
                page_size=page_size,
            )
        except Exception as e:
            logger.error(f"Oracle query execution failed: {e}")
            raise

    def _build_oracle_sql(
        self,
        qd: QueryDefinition,
        scope: DataScope,
        page: int,
        page_size: int,
    ) -> tuple[str, dict]:
        table = quote_table(qd.schema_name, qd.object_name)
        select_parts = []
        params: dict = {}

        for dim in qd.dimensions:
            select_parts.append(quote_ident(dim.field_name))
        for metric in qd.metrics:
            agg = safe_agg(metric.aggregation)
            field = quote_ident(metric.field_name)
            alias = quote_ident(metric.field_name)
            if agg == "COUNT_DISTINCT":
                select_parts.append(f"COUNT(DISTINCT {field}) {alias}")
            else:
                select_parts.append(f"{agg}({field}) {alias}")
        if not select_parts:
            select_parts = ["*"]

        select_clause = ", ".join(select_parts)
        where_clauses = apply_org_scope_clauses(qd, scope, params, prefix="ora_")

        where_sql = f"WHERE {' AND '.join(where_clauses)}" if where_clauses else ""

        group_by = ""
        if qd.dimensions and qd.metrics:
            dims = ", ".join(quote_ident(d.field_name) for d in qd.dimensions)
            group_by = f"GROUP BY {dims}"

        order_by = ""
        if qd.sort_by:
            direction = "DESC" if qd.sort_order.upper() == "DESC" else "ASC"
            order_by = f"ORDER BY {quote_ident(qd.sort_by)} {direction}"

        offset = (page - 1) * page_size
        sql = (
            f"SELECT {select_clause} FROM {table} "
            f"{where_sql} {group_by} {order_by} "
            f"OFFSET :offset ROWS FETCH NEXT :fetch_rows ROWS ONLY"
        )
        params["offset"] = offset
        params["fetch_rows"] = page_size
        return sql.strip(), params


# ── Internal API connector ────────────────────────────────────────────────────

class InternalAPIConnector(BaseConnector):
    """
    Fetches data from an internal REST API endpoint.
    Credentials remain on the backend; never exposed to the frontend.
    """

    def __init__(self, base_url: str, auth_config: dict):
        self.base_url = base_url
        self.auth_config = auth_config  # {"type": "bearer", "token": "..."} etc.

    def _get_headers(self) -> dict:
        auth_type = self.auth_config.get("type", "none")
        if auth_type == "bearer":
            return {"Authorization": f"Bearer {self.auth_config['token']}"}
        elif auth_type == "basic":
            import base64
            creds = base64.b64encode(
                f"{self.auth_config['username']}:{self.auth_config['password']}".encode()
            ).decode()
            return {"Authorization": f"Basic {creds}"}
        return {}

    def test_connection(self) -> bool:
        try:
            import httpx
            resp = httpx.get(self.base_url, headers=self._get_headers(), timeout=5)
            return resp.status_code < 500
        except Exception as e:
            logger.error(f"Internal API connection test failed: {e}")
            return False

    def get_schema_objects(self, schema: str) -> list[str]:
        # For API sources, endpoints are configured in the dataset
        return []

    def get_field_metadata(self, schema: str, object_name: str) -> list[dict]:
        # Field metadata is manually defined for API sources
        return []

    def execute_query(
        self,
        query_def: QueryDefinition,
        scope: DataScope,
        page: int = 1,
        page_size: int = 25,
    ) -> QueryResult:
        """
        Calls the configured endpoint, injects org-scope params as query params.
        Response is expected to be JSON array or {data: [...], total: N}.
        """
        if scope.deny_all:
            return empty_result(page, page_size)
        import httpx
        params: dict = {"page": page, "page_size": page_size}

        # Inject scope as query parameters the internal API understands
        if not scope.is_unrestricted:
            if scope.branch_ids:
                params["branch_ids"] = ",".join(str(b) for b in scope.branch_ids)
            elif scope.district_ids:
                params["district_ids"] = ",".join(str(d) for d in scope.district_ids)
            elif scope.region_ids:
                params["region_ids"] = ",".join(str(r) for r in scope.region_ids)
            else:
                return empty_result(page, page_size)

        endpoint = query_def.object_name  # endpoint path for API sources
        url = f"{self.base_url.rstrip('/')}/{endpoint.lstrip('/')}"

        try:
            resp = httpx.get(url, headers=self._get_headers(), params=params, timeout=30)
            resp.raise_for_status()
            data = resp.json()

            if isinstance(data, list):
                rows = data
                total = len(data)
            else:
                rows = data.get("data", [])
                total = data.get("total", len(rows))

            columns = list(rows[0].keys()) if rows else []
            return QueryResult(columns=columns, rows=rows, total_count=total, page=page, page_size=page_size)
        except Exception as e:
            logger.error(f"Internal API query failed: {e}")
            raise


class ExcelConnector(BaseConnector):
    """Read-only Excel workbook stored on the application server."""

    def __init__(self, file_path: str):
        self.file_path = file_path

    def test_connection(self) -> bool:
        import os
        return bool(self.file_path) and os.path.isfile(self.file_path)

    def get_schema_objects(self, schema: str) -> list[str]:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(self.file_path, read_only=True, data_only=True)
            names = list(wb.sheetnames)
            wb.close()
            return names
        except Exception as e:
            logger.error(f"Excel get_schema_objects failed: {e}")
            return []

    def get_field_metadata(self, schema: str, object_name: str) -> list[dict]:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(self.file_path, read_only=True, data_only=True)
            ws = wb[object_name] if object_name in wb.sheetnames else wb.active
            header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
            wb.close()
            return [
                {"field_name": str(h).strip(), "data_type": "TEXT", "nullable": True}
                for h in header if h
            ]
        except Exception as e:
            logger.error(f"Excel get_field_metadata failed: {e}")
            return []

    def _load_rows(self, sheet_name: Optional[str]) -> tuple[list[str], list[dict]]:
        import openpyxl
        wb = openpyxl.load_workbook(self.file_path, read_only=True, data_only=True)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        rows_iter = ws.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        if not header_row:
            wb.close()
            return [], []
        columns = [str(h).strip() for h in header_row if h is not None]
        data = []
        for raw in rows_iter:
            if all(v is None for v in raw):
                continue
            data.append({columns[i]: raw[i] if i < len(raw) else None for i in range(len(columns))})
        wb.close()
        return columns, data

    def execute_query(
        self,
        query_def: QueryDefinition,
        scope: DataScope,
        page: int = 1,
        page_size: int = 25,
    ) -> QueryResult:
        if scope.deny_all:
            return empty_result(page, page_size)
        columns, rows = self._load_rows(query_def.object_name)
        if not scope.is_unrestricted:
            col = None
            allowed = None
            if scope.branch_ids and query_def.branch_column:
                col, allowed = query_def.branch_column, set(scope.branch_ids)
            elif scope.district_ids and query_def.district_column:
                col, allowed = query_def.district_column, set(scope.district_ids)
            elif scope.region_ids and query_def.region_column:
                col, allowed = query_def.region_column, set(scope.region_ids)
            if not col:
                return empty_result(page, page_size)
            filtered = []
            for row in rows:
                try:
                    val = int(row.get(col))
                except (TypeError, ValueError):
                    continue
                if val in allowed:
                    filtered.append(row)
            rows = filtered
        for f in query_def.filters:
            field, op, val = f.get("field"), f.get("operator", "eq"), f.get("value")
            if not field:
                continue
            def _keep(row, field=field, op=op, val=val):
                cell = row.get(field)
                if op == "eq":
                    return str(cell) == str(val)
                if op == "gte":
                    return cell is not None and str(cell) >= str(val)
                if op == "lte":
                    return cell is not None and str(cell) <= str(val)
                return True
            rows = [r for r in rows if _keep(r)]
        if query_def.dimensions and not query_def.metrics:
            keep = [d.field_name for d in query_def.dimensions if d.field_name in (columns or [])]
            if keep:
                rows = [{k: r.get(k) for k in keep} for r in rows]
                columns = keep
        total = len(rows)
        start = max((page - 1) * page_size, 0)
        page_rows = rows[start:start + page_size]
        return QueryResult(
            columns=list(page_rows[0].keys()) if page_rows else columns,
            rows=page_rows,
            total_count=total,
            page=page,
            page_size=page_size,
        )


# ── Factory ───────────────────────────────────────────────────────────────────

def get_connector(source_type: str, config: dict) -> BaseConnector:
    """
    Factory function — returns the correct connector for a data source type.
    config dict contains decrypted connection parameters.
    """
    if source_type == "POSTGRESQL":
        return PostgreSQLConnector(
            host=config["host"],
            port=int(config.get("port", 5432)),
            database=config["database"],
            username=config["username"],
            password=config["password"],
            schema=config.get("schema", "public"),
        )
    elif source_type == "ORACLE":
        return OracleConnector(
            host=config["host"],
            port=int(config.get("port", 1521)),
            service_name=config["service_name"],
            username=config["username"],
            password=config["password"],
        )
    elif source_type == "INTERNAL_API":
        return InternalAPIConnector(
            base_url=config["api_url"],
            auth_config=config.get("auth_config", {}),
        )
    elif source_type == "EXCEL":
        return ExcelConnector(file_path=config.get("file_path") or config.get("api_url") or "")
    else:
        raise ValueError(f"Unknown data source type: {source_type}")
